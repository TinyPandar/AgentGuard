#!/usr/bin/env python3
"""Markdown -> DOCX for the competition report.

Applies the official附件2 formatting constraints:
  - A4 page
  - body text: 宋体 (SimSun), 小四 = 12pt, 1.5 line spacing
  - headings: 宋体, bold, larger

Usage (does not touch any project lockfile):
  uv run --no-project --with python-docx python3 build_report_docx.py <input.md> <output.docx>
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK = "宋体"
MONO = "Consolas"
BODY_PT = Pt(12)  # 小四
HEADING_PT = {1: Pt(16), 2: Pt(15), 3: Pt(14), 4: Pt(12.5)}

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`(.+?)`")
TABLE_SEP_RE = re.compile(r"^\|?[\s:|-]+\|?$")


def set_font(run, name: str = CJK, size: Pt | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold


def add_runs(paragraph, text: str, base_size: Pt = BODY_PT) -> None:
    """Render inline **bold** and `code`."""
    tokens = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for token in tokens:
        if not token:
            continue
        m = BOLD_RE.fullmatch(token)
        if m:
            set_font(paragraph.add_run(m.group(1)), CJK, base_size, True)
            continue
        m = CODE_RE.fullmatch(token)
        if m:
            set_font(paragraph.add_run(m.group(1)), MONO, Pt(base_size.pt - 1.5))
            continue
        set_font(paragraph.add_run(token), CJK, base_size)


def style_body(paragraph) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(24)  # 2 chars at 12pt


def strip_comments(lines: list[str]) -> list[str]:
    out, in_comment = [], False
    for line in lines:
        if in_comment:
            if "-->" in line:
                in_comment = False
            continue
        if line.lstrip().startswith("<!--"):
            if "-->" not in line:
                in_comment = True
            continue
        out.append(line)
    return out


def build(md_path: Path, docx_path: Path) -> None:
    lines = strip_comments(md_path.read_text("utf-8").splitlines())
    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(3.17))
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Cm(2.54))

    normal = doc.styles["Normal"]
    normal.font.name = CJK
    normal.font.size = BODY_PT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        # code fence
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Pt(18)
            for j, code_line in enumerate(block):
                if j:
                    p.add_run().add_break()
                set_font(p.add_run(code_line), MONO, Pt(9))
            continue

        # table
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [
                [c.strip() for c in r.strip("|").split("|")]
                for r in block
                if not TABLE_SEP_RE.fullmatch(r.replace(" ", ""))
            ]
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    if c >= len(table.columns):
                        continue
                    para = table.cell(r, c).paragraphs[0]
                    para.paragraph_format.line_spacing = 1.0
                    add_runs(para, cell, Pt(10.5))
                    if r == 0:
                        for run in para.runs:
                            run.font.bold = True
            doc.add_paragraph()
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(10 if level > 1 else 0)
            p.paragraph_format.space_after = Pt(4)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, m.group(2), HEADING_PT[level])
            for run in p.runs:
                run.font.bold = True
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = Pt(18)
            add_runs(p, stripped.lstrip("> ").strip(), Pt(10.5))
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # lists
        m = re.match(r"^([-*]|\d+\.)\s+(.*)$", stripped)
        if m:
            marker = "" if m.group(1) in ("-", "*") else m.group(1) + " "
            p = doc.add_paragraph(style="List Bullet" if marker == "" else "List Number")
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, marker + m.group(2))
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        style_body(p)
        add_runs(p, stripped)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"OK -> {docx_path} ({docx_path.stat().st_size} bytes)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(Path(sys.argv[1]), Path(sys.argv[2]))
